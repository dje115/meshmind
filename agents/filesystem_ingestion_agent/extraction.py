"""
Local-only extraction pipeline for MeshMind ingestion agent.
Direct reading first; OCR for scanned PDFs. No cloud.
"""

import hashlib
import os
import subprocess
import tempfile
from abc import ABC, abstractmethod
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from contract_models import IngestedChunk, IngestedItem, PIPELINE_VERSION

# Optional imports for extraction
try:
    import fitz  # PyMuPDF
    HAS_PYMUPDF = True
except ImportError:
    HAS_PYMUPDF = False

try:
    from docx import Document as DocxDocument
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    import pytesseract
    from PIL import Image
    HAS_OCR = True
except ImportError:
    HAS_OCR = False


@dataclass
class ExtractResult:
    text: str
    content_type: str
    extraction_method: str
    ocr_attempted: bool
    ocr_used: bool
    page_count: int
    failure_reason: str | None
    metadata: dict[str, Any]


def compute_file_hash(path: Path) -> str:
    """SHA-256 of file content."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def chunk_text(text: str, chunk_size: int = 1200, overlap: int = 200) -> list[IngestedChunk]:
    """Split text into overlapping chunks."""
    if not text.strip():
        return []
    chunks: list[IngestedChunk] = []
    start = 0
    idx = 0
    while start < len(text):
        end = start + chunk_size
        chunk_text_str = text[start:end]
        if chunk_text_str.strip():
            chunks.append(IngestedChunk(chunk_index=idx, chunk_text=chunk_text_str, page_number=0))
            idx += 1
        start = end - overlap
        if start >= len(text):
            break
    return chunks


class ExtractionProvider(ABC):
    """Abstract extraction provider."""

    @abstractmethod
    def supports(self, path: Path) -> bool:
        pass

    @abstractmethod
    def extract(self, path: Path) -> ExtractResult:
        pass


class PdfProvider(ExtractionProvider):
    """PDF extraction via PyMuPDF. OCR fallback for scanned PDFs."""

    def supports(self, path: Path) -> bool:
        return path.suffix.lower() == ".pdf"

    def extract(self, path: Path) -> ExtractResult:
        if not HAS_PYMUPDF:
            return ExtractResult(
                text="",
                content_type="application/pdf",
                extraction_method="none",
                ocr_attempted=False,
                ocr_used=False,
                page_count=0,
                failure_reason="PyMuPDF not installed",
                metadata={},
            )
        try:
            doc = fitz.open(path)
            page_count = len(doc)
            text_parts: list[str] = []
            for i in range(page_count):
                page = doc[i]
                t = page.get_text()
                text_parts.append(t)
            doc.close()
            text = "\n\n".join(text_parts)

            # OCR fallback if very little text
            ocr_attempted = False
            ocr_used = False
            if len(text.strip()) < 200 and HAS_OCR:
                ocr_result = _run_pdf_ocr(path)
                if ocr_result:
                    text = ocr_result
                    ocr_attempted = True
                    ocr_used = True

            return ExtractResult(
                text=text,
                content_type="application/pdf",
                extraction_method="pymupdf" + ("+ocr" if ocr_used else ""),
                ocr_attempted=ocr_attempted,
                ocr_used=ocr_used,
                page_count=page_count,
                failure_reason=None if text.strip() else "No extractable text",
                metadata={"page_count": page_count},
            )
        except Exception as e:
            return ExtractResult(
                text="",
                content_type="application/pdf",
                extraction_method="pymupdf",
                ocr_attempted=False,
                ocr_used=False,
                page_count=0,
                failure_reason=str(e),
                metadata={},
            )


def _run_pdf_ocr(path: Path) -> str | None:
    """Run pdftoppm + tesseract for scanned PDFs. Returns extracted text or None."""
    if not HAS_OCR:
        return None
    try:
        # Check for pdftoppm
        subprocess.run(["pdftoppm", "-v"], capture_output=True, check=False)
    except FileNotFoundError:
        return None
    try:
        with tempfile.TemporaryDirectory() as tmp:
            out_prefix = str(Path(tmp) / "page")
            subprocess.run(
                ["pdftoppm", "-png", "-r", "300", str(path), out_prefix],
                capture_output=True,
                check=True,
            )
            texts: list[str] = []
            for i in range(100):  # reasonable max pages
                img_path = Path(f"{out_prefix}-{i+1:06d}.png")
                if not img_path.exists():
                    break
                img = Image.open(img_path)
                t = pytesseract.image_to_string(img)
                texts.append(t)
            return "\n\n".join(texts).strip() if texts else None
    except Exception:
        return None


class DocxProvider(ExtractionProvider):
    """DOCX extraction via python-docx."""

    def supports(self, path: Path) -> bool:
        return path.suffix.lower() == ".docx"

    def extract(self, path: Path) -> ExtractResult:
        if not HAS_DOCX:
            return ExtractResult(
                text="",
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                extraction_method="none",
                ocr_attempted=False,
                ocr_used=False,
                page_count=0,
                failure_reason="python-docx not installed",
                metadata={},
            )
        try:
            doc = DocxDocument(path)
            paras = [p.text for p in doc.paragraphs if p.text.strip()]
            text = "\n".join(paras)
            return ExtractResult(
                text=text,
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                extraction_method="python-docx",
                ocr_attempted=False,
                ocr_used=False,
                page_count=0,
                failure_reason=None if text.strip() else "Document empty",
                metadata={},
            )
        except Exception as e:
            return ExtractResult(
                text="",
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                extraction_method="python-docx",
                ocr_attempted=False,
                ocr_used=False,
                page_count=0,
                failure_reason=str(e),
                metadata={},
            )


class XlsxProvider(ExtractionProvider):
    """XLSX extraction via openpyxl."""

    def supports(self, path: Path) -> bool:
        return path.suffix.lower() in (".xlsx", ".xlsm")

    def extract(self, path: Path) -> ExtractResult:
        if not HAS_OPENPYXL:
            return ExtractResult(
                text="",
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                extraction_method="none",
                ocr_attempted=False,
                ocr_used=False,
                page_count=0,
                failure_reason="openpyxl not installed",
                metadata={},
            )
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            parts: list[str] = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    cells = [str(c) if c is not None else "" for c in row]
                    if any(cells):
                        parts.append("\t".join(cells))
            wb.close()
            text = "\n".join(parts)
            return ExtractResult(
                text=text,
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                extraction_method="openpyxl",
                ocr_attempted=False,
                ocr_used=False,
                page_count=0,
                failure_reason=None if text.strip() else "Spreadsheet empty",
                metadata={},
            )
        except Exception as e:
            return ExtractResult(
                text="",
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                extraction_method="openpyxl",
                ocr_attempted=False,
                ocr_used=False,
                page_count=0,
                failure_reason=str(e),
                metadata={},
            )


class TextProvider(ExtractionProvider):
    """Plain text / markdown."""

    def supports(self, path: Path) -> bool:
        return path.suffix.lower() in (".txt", ".md", ".rtf")

    def extract(self, path: Path) -> ExtractResult:
        try:
            with open(path, encoding="utf-8", errors="replace") as f:
                text = f.read()
            ext = path.suffix.lower()
            ct = "text/plain" if ext in (".txt", ".md") else "application/rtf"
            return ExtractResult(
                text=text,
                content_type=ct,
                extraction_method="direct",
                ocr_attempted=False,
                ocr_used=False,
                page_count=0,
                failure_reason=None if text.strip() else "File empty",
                metadata={},
            )
        except Exception as e:
            return ExtractResult(
                text="",
                content_type="text/plain",
                extraction_method="direct",
                ocr_attempted=False,
                ocr_used=False,
                page_count=0,
                failure_reason=str(e),
                metadata={},
            )


DEFAULT_PROVIDERS: list[ExtractionProvider] = [
    PdfProvider(),
    DocxProvider(),
    XlsxProvider(),
    TextProvider(),
]


def get_provider(path: Path, providers: list[ExtractionProvider] | None = None) -> ExtractionProvider | None:
    """Return first provider that supports the path."""
    for p in providers or DEFAULT_PROVIDERS:
        if p.supports(path):
            return p
    return None


def extract_and_normalize(
    path: Path,
    source_id: str,
    source_parent: str | None,
    providers: list[ExtractionProvider] | None = None,
) -> IngestedItem:
    """
    Extract content and normalize into IngestedItem.
    """
    path = path.resolve()
    provider = get_provider(path, providers)

    if provider is None:
        return _make_failed_item(
            path,
            source_id,
            source_parent,
            "skipped_unsupported",
            f"Format {path.suffix} not supported",
        )

    result = provider.extract(path)

    if result.failure_reason and not result.text.strip():
        status = "failed_ocr" if result.ocr_attempted else "failed_extraction"
        return _make_failed_item(path, source_id, source_parent, status, result.failure_reason)

    chunks = chunk_text(result.text)
    mtime_ms = int(os.path.getmtime(path) * 1000)
    import time
    ingested_at_ms = int(time.time() * 1000)
    content_hash = compute_file_hash(path)

    return IngestedItem(
        source_id=source_id,
        source_type="filesystem",
        item_id=str(path),
        source_display_name=path.name,
        source_origin_label=str(path),
        source_locator=str(path),
        source_open_target=f"file:///{path.as_posix()}",
        path_or_external_key=str(path),
        content_type=result.content_type,
        extracted_text=result.text[:100_000],  # cap
        chunks=[c for c in chunks[:500]],  # cap chunks
        ocr_attempted=result.ocr_attempted,
        ocr_used=result.ocr_used,
        extraction_method=result.extraction_method,
        ingest_status="ingested",
        source_modified_at=mtime_ms,
        ingested_at=ingested_at_ms,
        content_hash=content_hash,
        pipeline_version=PIPELINE_VERSION,
        source_parent=source_parent,
        metadata=result.metadata,
    )


def _make_failed_item(
    path: Path,
    source_id: str,
    source_parent: str | None,
    status: str,
    failure_reason: str,
) -> IngestedItem:
    path = path.resolve()
    mtime_ms = int(os.path.getmtime(path) * 1000) if path.exists() else 0
    import time
    ingested_at_ms = int(time.time() * 1000)
    return IngestedItem(
        source_id=source_id,
        source_type="filesystem",
        item_id=str(path),
        source_display_name=path.name,
        source_origin_label=str(path),
        source_locator=str(path),
        source_open_target=f"file:///{path.as_posix()}",
        path_or_external_key=str(path),
        content_type="application/octet-stream",
        extracted_text="",
        chunks=[],
        ocr_attempted=False,
        ocr_used=False,
        extraction_method="none",
        ingest_status=status,
        source_modified_at=mtime_ms,
        ingested_at=ingested_at_ms,
        content_hash="",
        pipeline_version=PIPELINE_VERSION,
        source_parent=source_parent,
        failure_reason=failure_reason,
    )
